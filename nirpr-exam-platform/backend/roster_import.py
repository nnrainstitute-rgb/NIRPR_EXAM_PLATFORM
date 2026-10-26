"""Read bounded delimited text and Excel candidate lists without changing accounts."""
import csv
import io
import zipfile
from fastapi import HTTPException

MAX_ROWS = 500


def parse_roster_file(content, filename):
    extension = (filename or "").lower().rsplit(".", 1)[-1]
    try:
        if extension in ("csv", "tsv", "txt"):
            text = content.decode("utf-16") if content.startswith((b"\xff\xfe",b"\xfe\xff")) else content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            values = csv.reader(io.StringIO(text), delimiter="\t") if extension=="tsv" else csv.reader(io.StringIO(text), dialect)
            return extract_rows(values)
        if extension == "xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                sheet = workbook.sheet_by_index(0)
                return extract_rows(sheet.row_values(index) for index in range(sheet.nrows))
            finally:
                workbook.release_resources()
        if extension in ("xlsx", "xlsm"):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(item.file_size for item in archive.infolist()) > 30*1024*1024:
                    raise HTTPException(status_code=413, detail="Excel file expands beyond the allowed size")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                return extract_rows(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        raise HTTPException(status_code=400, detail="Upload .csv, .tsv, .txt, .xls, .xlsx or .xlsm")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read this file. Use a valid CSV/TSV/text file or an Excel .xls/.xlsx/.xlsm workbook")


def extract_rows(values):
    iterator = iter(values)
    header = next(iterator, None)
    if not header:
        raise HTTPException(status_code=400, detail="The file is empty")
    headers = [str(cell or "").strip().lower().replace(" ", "_") for cell in header]
    email_columns = [index for index, name in enumerate(headers) if name in ("email", "email_address")]
    if len(email_columns) != 1:
        raise HTTPException(status_code=400, detail="Include exactly one email column in the first row")
    column = email_columns[0]
    rows = []
    for number, row in enumerate(iterator, start=2):
        if number > 10002:
            raise HTTPException(status_code=413, detail="Too many worksheet rows; remove unused rows")
        if not any(str(cell or "").strip() for cell in row):
            continue
        if len(rows) >= MAX_ROWS:
            raise HTTPException(status_code=413, detail="Upload no more than 500 candidates at once")
        rows.append({"row": number, "email": str(row[column] or "").strip().lower() if column < len(row) else ""})
    if not rows:
        raise HTTPException(status_code=400, detail="Add at least one candidate email below the header")
    return rows
